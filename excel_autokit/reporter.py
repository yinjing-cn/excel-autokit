"""
Report generator for excel_autokit.

Produces formatted Excel reports using Jinja2 templates, with support
for chart embedding via matplotlib, batch generation by group keys,
and automatic summary/subtotal rows.
"""

import io
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd
from jinja2 import BaseLoader, Environment, TemplateNotFound

from excel_autokit.models import ReportConfig
from excel_autokit.utils import (
    DataValidationError,
    TemplateError,
    ensure_path,
    get_logger,
    validate_columns,
)

logger = get_logger(__name__)


class ReportGenerator:
    """Excel report generator with Jinja2 templating.

    Generates formatted Excel workbooks from data, with optional
    chart embedding, batch splitting by group keys, and auto-summary.

    Usage::

        gen = ReportGenerator()
        gen.load_data(df)
        gen.register_template("summary", template_str)
        gen.generate("summary", output="report.xlsx")
    """

    def __init__(self) -> None:
        """Initialize the report generator."""
        self._env = Environment(loader=BaseLoader())
        self._templates: Dict[str, Any] = {}
        self._df: Optional[pd.DataFrame] = None
        self._metadata: Dict[str, Any] = {}

    def load_data(self, df: pd.DataFrame, metadata: Optional[Dict[str, Any]] = None) -> "ReportGenerator":
        """Load data for report generation.

        Args:
            df: Source DataFrame.
            metadata: Optional metadata dict available in templates.

        Returns:
            Self for chaining.
        """
        self._df = df.copy()
        self._metadata = metadata or {}
        logger.info("ReportGenerator loaded %d rows", len(self._df))
        return self

    def register_template(self, name: str, template_str: str) -> "ReportGenerator":
        """Register a Jinja2 template string.

        Args:
            name: Template name for later reference.
            template_str: Jinja2 template content.

        Returns:
            Self for chaining.
        """
        self._templates[name] = self._env.from_string(template_str)
        logger.debug("Template '%s' registered", name)
        return self

    def register_template_file(self, name: str, template_path: Union[str, Path]) -> "ReportGenerator":
        """Register a template from a file.

        Args:
            name: Template name.
            template_path: Path to the template file.

        Returns:
            Self for chaining.

        Raises:
            TemplateError: If the template file cannot be read.
        """
        p = Path(template_path)
        if not p.exists():
            raise TemplateError(f"Template file not found: {p}")
        try:
            content = p.read_text(encoding="utf-8")
            return self.register_template(name, content)
        except Exception as exc:
            raise TemplateError(f"Failed to read template {p}: {exc}") from exc

    @property
    def template_names(self) -> List[str]:
        """Return list of registered template names."""
        return list(self._templates.keys())

    def generate(
        self,
        template_name: str,
        output: Union[str, Path] = "report.xlsx",
        group_by: Optional[List[str]] = None,
        include_summary: bool = True,
        chart_columns: Optional[List[str]] = None,
    ) -> Union[Path, List[Path]]:
        """Generate report(s) from data and template.

        Args:
            template_name: Name of the registered template.
            output: Output file path (or directory for batch).
            group_by: If set, generate one report per group.
            include_summary: Whether to add summary/subtotal rows.
            chart_columns: Columns to visualize as embedded charts.

        Returns:
            Path or list of paths to generated files.

        Raises:
            TemplateError: If template not found or rendering fails.
            DataValidationError: If no data is loaded.
        """
        if self._df is None:
            raise DataValidationError("No data loaded. Call load_data() first.")
        if template_name not in self._templates:
            raise TemplateError(f"Template '{template_name}' not registered")

        if group_by:
            return self._generate_batch(template_name, output, group_by, include_summary)
        return self._generate_single(template_name, output, include_summary, chart_columns)

    def _generate_single(
        self,
        template_name: str,
        output: Union[str, Path],
        include_summary: bool,
        chart_columns: Optional[List[str]],
    ) -> Path:
        """Generate a single report file."""
        df = self._df.copy()

        # Build template context
        context = self._build_context(df)

        # Render template for sheet names / formatting hints
        try:
            rendered = self._templates[template_name].render(context)
        except Exception as exc:
            raise TemplateError(f"Template rendering failed: {exc}") from exc

        path = ensure_path(output)

        with pd.ExcelWriter(str(path), engine="openpyxl") as writer:
            # Write main data sheet
            sheet_name = rendered.strip() if rendered.strip() and len(rendered.strip()) <= 31 else "Report"
            sheet_name = self._sanitize_sheet_name(sheet_name)
            df.to_excel(writer, sheet_name=sheet_name, index=False)

            # Add summary rows
            if include_summary:
                self._add_summary_rows(writer, sheet_name, df)

            # Add charts
            if chart_columns:
                self._add_chart_sheet(writer, df, chart_columns)

        logger.info("Report generated: %s", path)
        return path

    def _generate_batch(
        self,
        template_name: str,
        output: Union[str, Path],
        group_by: List[str],
        include_summary: bool,
    ) -> List[Path]:
        """Generate multiple reports, one per group."""
        output_dir = ensure_path(output)
        output_dir.mkdir(parents=True, exist_ok=True)

        validate_columns(self._df, group_by, "group_by")
        paths: List[Path] = []

        for group_values, group_df in self._df.groupby(group_by):
            # Build output filename
            if isinstance(group_values, tuple):
                suffix = "_".join(str(v) for v in group_values)
            else:
                suffix = str(group_values)
            filename = f"{template_name}_{suffix}.xlsx"
            file_path = output_dir / filename

            # Generate individual report
            context = self._build_context(group_df)
            rendered = self._templates[template_name].render(context)
            sheet_name = self._sanitize_sheet_name(rendered.strip() if rendered.strip() else "Report")

            with pd.ExcelWriter(str(file_path), engine="openpyxl") as writer:
                group_df.to_excel(writer, sheet_name=sheet_name, index=False)
                if include_summary:
                    self._add_summary_rows(writer, sheet_name, group_df)

            paths.append(file_path)
            logger.info("Batch report: %s (%d rows)", file_path.name, len(group_df))

        logger.info("Batch generation complete: %d files", len(paths))
        return paths

    def _build_context(self, df: pd.DataFrame) -> Dict[str, Any]:
        """Build the Jinja2 template rendering context."""
        context: Dict[str, Any] = {
            "data": df,
            "rows": df.to_dict("records"),
            "columns": list(df.columns),
            "row_count": len(df),
            "metadata": self._metadata,
        }
        # Add numeric summaries for each numeric column
        numeric_cols = df.select_dtypes(include="number").columns
        for col in numeric_cols:
            context[f"{col}_sum"] = df[col].sum()
            context[f"{col}_mean"] = round(df[col].mean(), 4) if len(df) > 0 else 0
            context[f"{col}_min"] = df[col].min()
            context[f"{col}_max"] = df[col].max()
        return context

    def _add_summary_rows(self, writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
        """Append summary/subtotal rows to the worksheet."""
        ws = writer.sheets[sheet_name]
        numeric_cols = df.select_dtypes(include="number").columns

        if numeric_cols.empty:
            return

        summary_row = ws.max_row + 2
        ws.cell(row=summary_row, column=1, value="Summary")

        for col_idx, col_name in enumerate(df.columns, start=1):
            if col_name in numeric_cols:
                total = df[col_name].sum()
                ws.cell(row=summary_row + 1, column=col_idx, value=f"Total: {total}")
            else:
                ws.cell(row=summary_row + 1, column=col_idx, value="")

    def _add_chart_sheet(
        self,
        writer: pd.ExcelWriter,
        df: pd.DataFrame,
        chart_columns: List[str],
    ) -> None:
        """Add a chart sheet with matplotlib-generated charts embedded as images."""
        try:
            import matplotlib
            matplotlib.use("Agg")
            import matplotlib.pyplot as plt
        except ImportError:
            logger.warning("matplotlib not available, skipping chart generation")
            return

        valid_cols = [c for c in chart_columns if c in df.columns]
        if not valid_cols:
            logger.warning("No valid chart columns found: %s", chart_columns)
            return

        numeric_df = df[valid_cols].select_dtypes(include="number")
        if numeric_df.empty:
            return

        fig, axes = plt.subplots(1, len(numeric_df.columns), figsize=(6 * len(numeric_df.columns), 4))
        if len(numeric_df.columns) == 1:
            axes = [axes]

        for ax, col in zip(axes, numeric_df.columns):
            numeric_df[col].plot(kind="bar", ax=ax, title=col)
            ax.tick_params(axis="x", rotation=45)

        plt.tight_layout()

        # Save chart to buffer and insert into Excel
        buf = io.BytesIO()
        fig.savefig(buf, format="png", dpi=100, bbox_inches="tight")
        buf.seek(0)
        plt.close(fig)

        # Write chart to a separate sheet using openpyxl image
        from openpyxl.drawing.image import Image as XlImage

        chart_ws = writer.book.create_sheet("Charts")
        img = XlImage(buf)
        chart_ws.add_image(img, "A1")

    @staticmethod
    def _sanitize_sheet_name(name: str) -> str:
        """Sanitize a string for use as an Excel sheet name."""
        # Excel sheet names: max 31 chars, no special chars
        invalid = r'\/:*?[]'
        for ch in invalid:
            name = name.replace(ch, "_")
        return name[:31] if name else "Sheet1"

# Created: 2026-09-08
