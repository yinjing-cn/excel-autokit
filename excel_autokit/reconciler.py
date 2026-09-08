"""
Reconciliation engine for excel_autokit.

Reads two Excel sources, performs composite-key matching with optional
numeric tolerance, detects differences, and outputs a styled result
Excel file with highlighted discrepancies.
"""

import logging
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill

from excel_autokit.models import DiffDetail, MatchKey, ReconcileResult, ReconcileSummary, Tolerance
from excel_autokit.utils import (
    DataValidationError,
    ReconciliationError,
    ensure_path,
    get_logger,
    read_excel_safe,
    validate_columns,
)

logger = get_logger(__name__)

# Styling constants
_FILL_RED = PatternFill(start_color="FFCCCC", end_color="FFCCCC", fill_type="solid")
_FILL_YELLOW = PatternFill(start_color="FFFFCC", end_color="FFFFCC", fill_type="solid")
_FONT_BOLD = Font(bold=True)


class Reconciler:
    """Excel reconciliation engine.

    Compares two data sources by composite keys, detects value differences
    with configurable numeric tolerance, and produces a styled output
    workbook highlighting discrepancies.

    Usage::

        reconciler = Reconciler(match_keys=["order_id", "sku"])
        summary = reconciler.reconcile("source_a.xlsx", "source_b.xlsx")
        reconciler.write_result("reconcile_output.xlsx")
    """

    def __init__(
        self,
        match_keys: Optional[List[str]] = None,
        tolerance: Optional[Tolerance] = None,
        compare_columns: Optional[List[Dict[str, Any]]] = None,
        ignore_case: bool = False,
    ):
        """Initialize the reconciler.

        Args:
            match_keys: Column names forming the composite match key.
            tolerance: Numeric tolerance for fuzzy matching.
            compare_columns: Column-level comparison config with per-column tolerance.
            ignore_case: Whether to ignore case in string key matching.
        """
        self._match_key = MatchKey(columns=match_keys or [], ignore_case=ignore_case)
        self._tolerance = tolerance or Tolerance()
        self._compare_columns = compare_columns or []
        self._df_a: Optional[pd.DataFrame] = None
        self._df_b: Optional[pd.DataFrame] = None
        self._summary: Optional[ReconcileSummary] = None
        self._result_df: Optional[pd.DataFrame] = None

    @classmethod
    def from_rule_engine(cls, rule_engine: Any) -> "Reconciler":
        """Create a Reconciler from a RuleEngine instance.

        Args:
            rule_engine: A configured RuleEngine.

        Returns:
            Reconciler configured per the rules.
        """
        return cls(
            match_keys=rule_engine.match_key.columns,
            tolerance=rule_engine.tolerance,
            compare_columns=rule_engine.compare_columns,
            ignore_case=rule_engine.match_key.ignore_case,
        )

    def reconcile(
        self,
        source_a: Union[str, Path, pd.DataFrame],
        source_b: Union[str, Path, pd.DataFrame],
        sheet_a: Union[str, int] = 0,
        sheet_b: Union[str, int] = 0,
    ) -> ReconcileSummary:
        """Run reconciliation between two sources.

        Args:
            source_a: First data source (file path or DataFrame).
            source_b: Second data source (file path or DataFrame).
            sheet_a: Sheet name/index for source A (if file path).
            sheet_b: Sheet name/index for source B (if file path).

        Returns:
            ReconcileSummary with match statistics and details.

        Raises:
            ReconciliationError: If match keys are not configured.
            DataValidationError: If required columns are missing.
        """
        if not self._match_key.columns:
            raise ReconciliationError("Match keys must be configured before reconciliation")

        # Load data
        self._df_a = self._load_source(source_a, sheet_a, "source_a")
        self._df_b = self._load_source(source_b, sheet_b, "source_b")

        # Validate columns
        validate_columns(self._df_a, self._match_key.columns, "source_a")
        validate_columns(self._df_b, self._match_key.columns, "source_b")

        logger.info(
            "Reconciling: source_a=%d rows, source_b=%d rows, keys=%s",
            len(self._df_a), len(self._df_b), self._match_key.columns,
        )

        # Build lookup maps
        map_a = self._build_key_map(self._df_a)
        map_b = self._build_key_map(self._df_b)

        # Perform matching
        results = self._perform_matching(map_a, map_b)

        # Build summary
        self._summary = self._build_summary(results)
        logger.info(
            "Reconciliation complete: matched=%d, mismatched=%d, missing_a=%d, missing_b=%d",
            self._summary.matched, self._summary.mismatched,
            self._summary.missing_in_a, self._summary.missing_in_b,
        )

        # Build result DataFrame
        self._result_df = self._build_result_dataframe()

        return self._summary

    def _load_source(
        self, source: Union[str, Path, pd.DataFrame], sheet: Union[str, int], label: str
    ) -> pd.DataFrame:
        """Load a data source from file path or return DataFrame as-is."""
        if isinstance(source, pd.DataFrame):
            return source.copy()
        return read_excel_safe(source, sheet_name=sheet)

    def _build_key_map(self, df: pd.DataFrame) -> Dict[Tuple, List[int]]:
        """Build a mapping from composite key tuples to row indices."""
        key_map: Dict[Tuple, List[int]] = {}
        for idx, row in df.iterrows():
            key = self._match_key.extract(row.to_dict())
            key_map.setdefault(key, []).append(idx)
        return key_map

    def _perform_matching(
        self, map_a: Dict[Tuple, List[int]], map_b: Dict[Tuple, List[int]]
    ) -> List[ReconcileResult]:
        """Match records between two key maps and detect differences."""
        results: List[ReconcileResult] = []
        all_keys = set(map_a.keys()) | set(map_b.keys())

        for key in sorted(all_keys, key=str):
            in_a = key in map_a
            in_b = key in map_b

            if in_a and in_b:
                # Compare each pair (support 1:1 for simplicity)
                for idx_a, idx_b in zip(map_a[key], map_b[key]):
                    row_a = self._df_a.loc[idx_a].to_dict()
                    row_b = self._df_b.loc[idx_b].to_dict()
                    diffs = self._compare_rows(row_a, row_b)
                    status = "mismatched" if diffs else "matched"
                    results.append(ReconcileResult(
                        status=status, key_values=key,
                        row_a=row_a, row_b=row_b, differences=diffs,
                    ))
                # Handle unmatched extras if counts differ
                len_a, len_b = len(map_a[key]), len(map_b[key])
                if len_a > len_b:
                    for extra_idx in map_a[key][len_b:]:
                        row_a = self._df_a.loc[extra_idx].to_dict()
                        results.append(ReconcileResult(
                            status="missing_in_b", key_values=key,
                            row_a=row_a, row_b=None,
                        ))
                elif len_b > len_a:
                    for extra_idx in map_b[key][len_a:]:
                        row_b = self._df_b.loc[extra_idx].to_dict()
                        results.append(ReconcileResult(
                            status="missing_in_a", key_values=key,
                            row_a=None, row_b=row_b,
                        ))
            elif in_a:
                for idx_a in map_a[key]:
                    row_a = self._df_a.loc[idx_a].to_dict()
                    results.append(ReconcileResult(
                        status="missing_in_b", key_values=key,
                        row_a=row_a, row_b=None,
                    ))
            else:
                for idx_b in map_b[key]:
                    row_b = self._df_b.loc[idx_b].to_dict()
                    results.append(ReconcileResult(
                        status="missing_in_a", key_values=key,
                        row_a=None, row_b=row_b,
                    ))

        return results

    def _compare_rows(self, row_a: Dict[str, Any], row_b: Dict[str, Any]) -> List[DiffDetail]:
        """Compare two matched rows and return differences."""
        diffs: List[DiffDetail] = []
        columns_to_check = self._get_compare_column_names(row_a, row_b)

        for col in columns_to_check:
            val_a = row_a.get(col)
            val_b = row_b.get(col)

            if self._values_equal(val_a, val_b, col):
                continue

            diffs.append(DiffDetail(column=col, value_a=val_a, value_b=val_b))

        return diffs

    def _get_compare_column_names(self, row_a: Dict, row_b: Dict) -> List[str]:
        """Determine which columns to compare."""
        if self._compare_columns:
            return [c["name"] for c in self._compare_columns if c["name"] in row_a or c["name"] in row_b]
        # Compare all non-key columns
        key_set = set(self._match_key.columns)
        all_cols = (set(row_a.keys()) | set(row_b.keys())) - key_set
        return sorted(all_cols)

    def _values_equal(self, val_a: Any, val_b: Any, column: str = "") -> bool:
        """Check if two values are equal, considering tolerance for numerics."""
        # Handle NaN
        a_nan = pd.isna(val_a) if not isinstance(val_a, (int, str)) else False
        b_nan = pd.isna(val_b) if not isinstance(val_b, (int, str)) else False
        if a_nan and b_nan:
            return True
        if a_nan != b_nan:
            return False

        # Try numeric comparison with tolerance
        try:
            num_a, num_b = float(val_a), float(val_b)
            # Check per-column tolerance first
            col_tol = self._get_column_tolerance(column)
            if col_tol is not None:
                return abs(num_a - num_b) <= col_tol
            return self._tolerance.is_within(num_a, num_b)
        except (TypeError, ValueError):
            pass

        # Fall back to direct comparison
        return val_a == val_b

    def _get_column_tolerance(self, column: str) -> Optional[float]:
        """Get per-column tolerance if configured."""
        for cc in self._compare_columns:
            if cc.get("name") == column and "tolerance" in cc:
                return cc["tolerance"]
        return None

    def _build_summary(self, results: List[ReconcileResult]) -> ReconcileSummary:
        """Aggregate reconciliation results into a summary."""
        summary = ReconcileSummary(
            total_a=len(self._df_a) if self._df_a is not None else 0,
            total_b=len(self._df_b) if self._df_b is not None else 0,
            details=results,
        )
        total_diff = 0.0
        for r in results:
            if r.status == "matched":
                summary.matched += 1
            elif r.status == "mismatched":
                summary.mismatched += 1
                # Sum amount differences for numeric diff columns
                for d in r.differences:
                    try:
                        total_diff += abs(float(d.value_b) - float(d.value_a))
                    except (TypeError, ValueError):
                        pass
            elif r.status == "missing_in_a":
                summary.missing_in_a += 1
            elif r.status == "missing_in_b":
                summary.missing_in_b += 1

        summary.total_amount_diff = total_diff
        return summary

    def _build_result_dataframe(self) -> pd.DataFrame:
        """Convert reconciliation results into a flat DataFrame."""
        rows = []
        key_cols = self._match_key.columns

        for r in self._summary.details:
            base: Dict[str, Any] = {"_status": r.status}
            for i, col in enumerate(key_cols):
                base[col] = r.key_values[i] if i < len(r.key_values) else ""

            # Merge data from both sources
            all_data_cols = set()
            if r.row_a:
                all_data_cols.update(r.row_a.keys())
            if r.row_b:
                all_data_cols.update(r.row_b.keys())
            data_cols = sorted(all_data_cols - set(key_cols))

            for col in data_cols:
                val_a = r.row_a.get(col) if r.row_a else None
                val_b = r.row_b.get(col) if r.row_b else None
                base[f"{col}_a"] = val_a
                base[f"{col}_b"] = val_b
                if val_a != val_b and val_a is not None and val_b is not None:
                    base[f"{col}_diff"] = "*"
                else:
                    base[f"{col}_diff"] = ""

            rows.append(base)

        return pd.DataFrame(rows) if rows else pd.DataFrame()

    def write_result(self, output_path: Union[str, Path]) -> Path:
        """Write reconciliation result to a styled Excel file.

        Args:
            output_path: Output Excel file path.

        Returns:
            Resolved path of the written file.

        Raises:
            ReconciliationError: If no reconciliation has been run yet.
        """
        if self._result_df is None or self._summary is None:
            raise ReconciliationError("No reconciliation result available. Run reconcile() first.")

        path = ensure_path(output_path)

        with pd.ExcelWriter(str(path), engine="openpyxl") as writer:
            self._result_df.to_excel(writer, sheet_name="Results", index=False)
            self._write_summary_sheet(writer)

            # Apply styling
            ws = writer.sheets["Results"]
            self._apply_result_styling(ws)

        logger.info("Reconciliation result written to %s", path)
        return path

    def _write_summary_sheet(self, writer: pd.ExcelWriter) -> None:
        """Write summary statistics sheet."""
        summary = self._summary
        rows = [
            ("Total records (Source A)", summary.total_a),
            ("Total records (Source B)", summary.total_b),
            ("Matched", summary.matched),
            ("Mismatched", summary.mismatched),
            ("Missing in Source A", summary.missing_in_a),
            ("Missing in Source B", summary.missing_in_b),
            ("Total Amount Difference", round(summary.total_amount_diff, 4)),
        ]
        df = pd.DataFrame(rows, columns=["Metric", "Value"])
        df.to_excel(writer, sheet_name="Summary", index=False)

    def _apply_result_styling(self, ws: Any) -> None:
        """Apply highlight styling to the results worksheet."""
        status_col = 1  # Column A = _status
        for row_idx in range(2, ws.max_row + 1):
            status = ws.cell(row=row_idx, column=status_col).value
            if status in ("mismatched",):
                for col_idx in range(1, ws.max_column + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = _FILL_RED
            elif status in ("missing_in_a", "missing_in_b"):
                for col_idx in range(1, ws.max_column + 1):
                    ws.cell(row=row_idx, column=col_idx).fill = _FILL_YELLOW

    @property
    def summary(self) -> Optional[ReconcileSummary]:
        """Return the last reconciliation summary, if available."""
        return self._summary

    @property
    def result_dataframe(self) -> Optional[pd.DataFrame]:
        """Return the result DataFrame, if available."""
        return self._result_df

# Created: 2026-09-08
