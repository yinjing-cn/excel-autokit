"""
Tests for the reconciliation engine (reconciler.py).
"""

import pandas as pd
import pytest

from excel_autokit.models import Tolerance
from excel_autokit.reconciler import Reconciler
from excel_autokit.utils import ReconciliationError


class TestReconcilerInit:
    """Test Reconciler initialization."""

    def test_default_init(self):
        r = Reconciler()
        assert r.summary is None
        assert r.result_dataframe is None

    def test_init_with_keys(self):
        r = Reconciler(match_keys=["order_id"])
        assert r._match_key.columns == ["order_id"]

    def test_init_with_tolerance(self):
        tol = Tolerance(absolute=0.05)
        r = Reconciler(match_keys=["id"], tolerance=tol)
        assert r._tolerance.absolute == 0.05

    def test_from_rule_engine(self, sample_recon_config):
        from excel_autokit.rules import RuleEngine
        engine = RuleEngine(config_dict=sample_recon_config)
        r = Reconciler.from_rule_engine(engine)
        assert r._match_key.columns == ["order_id", "sku"]
        assert r._tolerance.absolute == 0.01


class TestReconcilerReconcile:
    """Test reconciliation logic."""

    def test_reconcile_basic(self, sample_order_data_a, sample_order_data_b):
        r = Reconciler(match_keys=["order_id", "sku"])
        summary = r.reconcile(sample_order_data_a, sample_order_data_b)

        assert summary.total_a == 5
        assert summary.total_b == 4
        assert summary.matched >= 1  # ORD001 is a perfect match
        assert summary.mismatched >= 1  # ORD002 has quantity/amount diffs

    def test_reconcile_missing_records(self, sample_order_data_a, sample_order_data_b):
        r = Reconciler(match_keys=["order_id", "sku"])
        summary = r.reconcile(sample_order_data_a, sample_order_data_b)

        # ORD004, ORD005 are only in A -> missing_in_b
        assert summary.missing_in_b == 2
        # ORD006 is only in B -> missing_in_a
        assert summary.missing_in_a == 1

    def test_reconcile_with_tolerance(self):
        df_a = pd.DataFrame({"id": [1, 2], "amount": [100.0, 200.0]})
        df_b = pd.DataFrame({"id": [1, 2], "amount": [100.005, 200.5]})

        r = Reconciler(match_keys=["id"], tolerance=Tolerance(absolute=0.01))
        summary = r.reconcile(df_a, df_b)

        # id=1: within tolerance -> matched; id=2: diff=0.5 -> mismatched
        assert summary.matched == 1
        assert summary.mismatched == 1

    def test_reconcile_no_keys_raises(self, sample_order_data_a, sample_order_data_b):
        r = Reconciler()
        with pytest.raises(ReconciliationError, match="Match keys must be configured"):
            r.reconcile(sample_order_data_a, sample_order_data_b)

    def test_reconcile_missing_columns_raises(self):
        df_a = pd.DataFrame({"id": [1], "val": [10]})
        df_b = pd.DataFrame({"other": [1], "val": [10]})

        r = Reconciler(match_keys=["id"])
        with pytest.raises(Exception):
            r.reconcile(df_a, df_b)

    def test_reconcile_from_files(self, excel_file_a, excel_file_b):
        r = Reconciler(match_keys=["order_id", "sku"])
        summary = r.reconcile(excel_file_a, excel_file_b)
        assert summary.total_a == 5
        assert summary.total_b == 4

    def test_result_dataframe_available(self, sample_order_data_a, sample_order_data_b):
        r = Reconciler(match_keys=["order_id", "sku"])
        r.reconcile(sample_order_data_a, sample_order_data_b)
        assert r.result_dataframe is not None
        assert len(r.result_dataframe) > 0

    def test_ignore_case_matching(self):
        df_a = pd.DataFrame({"id": ["abc"], "val": [1]})
        df_b = pd.DataFrame({"id": ["ABC"], "val": [1]})

        r = Reconciler(match_keys=["id"], ignore_case=True)
        summary = r.reconcile(df_a, df_b)
        assert summary.matched == 1

    def test_summary_details_populated(self, sample_order_data_a, sample_order_data_b):
        r = Reconciler(match_keys=["order_id", "sku"])
        summary = r.reconcile(sample_order_data_a, sample_order_data_b)
        assert len(summary.details) > 0
        for detail in summary.details:
            assert detail.status in ("matched", "mismatched", "missing_in_a", "missing_in_b")


class TestReconcilerOutput:
    """Test reconciliation output writing."""

    def test_write_result(self, sample_order_data_a, sample_order_data_b, tmp_dir):
        r = Reconciler(match_keys=["order_id", "sku"])
        r.reconcile(sample_order_data_a, sample_order_data_b)
        output = tmp_dir / "result.xlsx"
        path = r.write_result(output)
        assert path.exists()

        # Verify sheets
        xls = pd.ExcelFile(str(path))
        assert "Results" in xls.sheet_names
        assert "Summary" in xls.sheet_names

    def test_write_result_before_reconcile_raises(self, tmp_dir):
        r = Reconciler(match_keys=["id"])
        with pytest.raises(ReconciliationError, match="No reconciliation result"):
            r.write_result(tmp_dir / "out.xlsx")

    def test_summary_sheet_content(self, sample_order_data_a, sample_order_data_b, tmp_dir):
        r = Reconciler(match_keys=["order_id", "sku"])
        r.reconcile(sample_order_data_a, sample_order_data_b)
        output = tmp_dir / "result.xlsx"
        r.write_result(output)

        summary_df = pd.read_excel(str(output), sheet_name="Summary")
        assert len(summary_df) == 7  # 7 metrics
        assert "Metric" in summary_df.columns
        assert "Value" in summary_df.columns

    def test_styled_output(self, sample_order_data_a, sample_order_data_b, tmp_dir):
        """Verify that the output Excel has styling applied."""
        from openpyxl import load_workbook

        r = Reconciler(match_keys=["order_id", "sku"])
        r.reconcile(sample_order_data_a, sample_order_data_b)
        output = tmp_dir / "styled_result.xlsx"
        r.write_result(output)

        wb = load_workbook(str(output))
        ws = wb["Results"]
        # Check that some cells have fill applied (mismatched rows)
        has_fill = False
        for row in ws.iter_rows(min_row=2):
            if row[0].fill and row[0].fill.start_color and row[0].fill.start_color.rgb:
                if row[0].fill.start_color.rgb != "00000000":
                    has_fill = True
                    break
        assert has_fill, "Expected at least one highlighted row in styled output"

# Created: 2026-09-08
